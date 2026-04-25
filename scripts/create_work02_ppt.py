#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Work-02: 中国城镇供水排水协会2026年会 PPT
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BLUE = RGBColor(0, 126, 181)
DARK_GRAY = RGBColor(77, 77, 77)
LIGHT_GRAY = RGBColor(153, 153, 153)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

def rect(s, x, y, w, h, fill=None):
    shape = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    shape.line.fill.background()
    return shape

def txt(s, text, x, y, w, h, sz=14, bold=False, col=DARK_GRAY,
        align=PP_ALIGN.LEFT):
    txBox = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(sz)
    run.font.bold = bold
    run.font.color.rgb = col
    return txBox

def bullet_list(s, items, x, y, w, h, sz=13, col=DARK_GRAY):
    txBox = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = item
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(sz)
        run.font.color.rgb = col
    return txBox

# ─── Slide 1: 封面 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 0.12, 7.5, fill=BLUE)
txt(s, "中国城镇供水排水协会", 0.5, 1.3, 12, 0.8, sz=38, bold=True, col=DARK_GRAY)
txt(s, "2026年年会", 0.5, 2.1, 12, 0.9, sz=44, bold=True, col=BLUE)
txt(s, "市场情报速览", 0.5, 3.1, 12, 0.6, sz=24, col=LIGHT_GRAY)
rect(s, 0.5, 3.8, 12.5, 0.03, fill=LIGHT_GRAY)
txt(s, "数据来源：百度百科、中国新闻网、深圳政府在线、南方新闻网  |  2026-04-22",
    0.5, 4.0, 12, 0.4, sz=11, col=LIGHT_GRAY)
rect(s, 0.5, 4.6, 3.5, 0.55, fill=BLUE)
txt(s, "深圳国际会展中心（宝安新馆）", 0.6, 4.65, 3.4, 0.45, sz=12, col=WHITE)
txt(s, "Hach中国IMS产品线", 0.5, 5.3, 12, 0.4, sz=11, col=LIGHT_GRAY)

# ─── Slide 2: 基础信息 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "年会基础信息", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

info = [
    ("名称", "中国城镇供水排水协会2026年年会（简称中国水协2026年会）"),
    ("时间", "2026年4月15日-19日（共5天）"),
    ("地点", "深圳国际会展中心（宝安新馆），主会场17号馆"),
    ("规模", "近万名代表，300余位专家，180余家参展企业，5万平米展示"),
    ("指导单位", "住房和城乡建设部城市建设司"),
    ("主办单位", "中国城镇供水排水协会、广东省城镇供水协会、深圳市环境水务集团等"),
]
for i, (k, v) in enumerate(info):
    ry = 1.1 + i * 0.98
    bg = RGBColor(250,250,250) if i % 2 == 0 else RGBColor(245,248,250)
    rect(s, 0.5, ry, 12.33, 0.88, fill=bg)
    rect(s, 0.5, ry, 0.06, 0.88, fill=BLUE)
    txt(s, k, 0.7, ry+0.1, 2.0, 0.7, sz=13, bold=True, col=BLUE)
    txt(s, v, 2.8, ry+0.1, 9.8, 0.7, sz=13, col=DARK_GRAY)

# ─── Slide 3: 六大核心议题 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "六大核心议题", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

topics = [
    ("规划落实", "聚焦十五五规划贯彻落实", BLUE),
    ("条例实施", "《供水条例》立法本意与实施路径", DARK_GRAY),
    ("人工智能", "多模态AI智能体、水务大模型、人工智能+全业务链融合", BLUE),
    ("系统治理", "厂-网-河一体化城镇水污染防控", DARK_GRAY),
    ("韧性城市", "内涝防治体系（智能感知+精准预警+科学处置）", BLUE),
    ("价格机制", "科学定价保障行业可持续发展", DARK_GRAY),
]
for i, (title, desc, col) in enumerate(topics):
    ry = 1.05 + i * 1.02
    rect(s, 0.5, ry, 12.33, 0.92, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 12.33, 0.06, fill=col)
    txt(s, title, 0.65, ry+0.15, 2.2, 0.7, sz=15, bold=True, col=col)
    txt(s, desc, 2.9, ry+0.15, 9.7, 0.7, sz=13, col=DARK_GRAY)

# ─── Slide 4: 日程 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "年会日程", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

schedule = [
    ("4月14日", "新闻发布会"),
    ("4月15日", "会长办公会、省级地方水协会长工作会"),
    ("4月16日", "第三届十七次常务理事会、第三届十八次理事会"),
    ("4月17日", "综合大会（上午）、技术交流与圆桌对话会（下午）"),
    ("4月18日", "技术交流与圆桌对话会（全天）"),
    ("4月19日", "深圳市典型水务项目参观考察（智慧水厂、全地下水质净化厂等）"),
]
for i, (date, event) in enumerate(schedule):
    ry = 1.1 + i * 1.0
    bg = RGBColor(250,250,250) if i % 2 == 0 else RGBColor(245,248,250)
    rect(s, 0.5, ry, 12.33, 0.9, fill=bg)
    rect(s, 0.5, ry, 1.8, 0.9, fill=BLUE)
    txt(s, date, 0.55, ry+0.2, 1.7, 0.5, sz=13, bold=True, col=WHITE, align=PP_ALIGN.CENTER)
    txt(s, event, 2.4, ry+0.15, 10.2, 0.65, sz=13, col=DARK_GRAY)

# ─── Slide 5: 关键信息 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "关键信息", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

key_info = [
    "深圳环境水务集团开放11条专业参观线路，提供深圳样板",
    "年会同期举办城镇水务技术与产品展示，5万平米，180余家水务材料设备企业",
    "2026年是十五五规划开局之年",
    "现场有自助签到制证机、人形机器人、AI打卡装置等智慧化服务设施",
    "华为公共事业军团参与，展示水务鸿蒙化解决方案",
    "AI议题贯穿全场：水务大模型、多模态AI智能体、人工智能+全业务链融合",
]
for i, info_text in enumerate(key_info):
    ry = 1.1 + i * 0.95
    bg = RGBColor(250,250,250) if i % 2 == 0 else RGBColor(232,244,252)
    rect(s, 0.5, ry, 12.33, 0.85, fill=bg)
    rect(s, 0.5, ry, 0.06, 0.85, fill=BLUE)
    txt(s, "* " + info_text, 0.7, ry+0.15, 12, 0.6, sz=13, col=DARK_GRAY)

# ─── Slide 6: AI议题深度 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "AI议题：从概念到落地", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)
txt(s, "Hach IMS视角关联", 0.5, 0.9, 12, 0.4, sz=13, col=LIGHT_GRAY)

ai_data = [
    ("水务大模型",
     "深圳水务深水云脑2.0已落地，五指耙水厂实证",
     "AI分析平台是IMS数据变现的出口"),
    ("多模态AI智能体",
     "华为展示水务场景多模态推理能力",
     "IMS提供多参数仪表数据是AI训练的基础"),
    ("人工智能+全业务链",
     "水协推动AI与运营/维护/决策全链条融合",
     "IMS是设备层数据接入AI的最小闭环"),
    ("智能感知+精准预警",
     "华为PLC+边缘计算+AI联动",
     "IMS+仪表是感知层的核心数据源"),
]
for i, (topic, desc, ims_link) in enumerate(ai_data):
    ry = 1.5 + i * 1.4
    rect(s, 0.5, ry, 12.33, 1.25, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 0.06, 1.25, fill=BLUE)
    txt(s, topic, 0.7, ry+0.1, 3.0, 0.5, sz=14, bold=True, col=BLUE)
    txt(s, desc, 0.7, ry+0.55, 5.5, 0.65, sz=12, col=DARK_GRAY)
    rect(s, 6.3, ry+0.1, 6.53, 1.05, fill=RGBColor(232,244,252))
    txt(s, ims_link, 6.45, ry+0.3, 6.2, 0.7, sz=12, col=BLUE)

# ─── Slide 7: 同期活动 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "同期活动：城镇水务技术与产品展示", 0.5, 0.35, 12, 0.6, sz=26, bold=True, col=DARK_GRAY)

items = [
    ("规模", "5万平米，180余家水务材料设备企业"),
    ("时间", "2026年4月16日-18日"),
    ("地点", "深圳国际会展中心（与年会同期同地）"),
    ("特点", "智慧化服务设施：自助签到制证机、人形机器人、AI打卡装置"),
    ("Hach机会", "180余家参展商中含大量潜在CP/集成商partner"),
]
for i, (k, v) in enumerate(items):
    ry = 1.1 + i * 1.0
    bg = RGBColor(250,250,250) if i % 2 == 0 else RGBColor(245,248,250)
    col = BLUE if k == "Hach机会" else DARK_GRAY
    bg_col = RGBColor(232,244,252) if k == "Hach机会" else bg
    rect(s, 0.5, ry, 12.33, 0.9, fill=bg_col)
    rect(s, 0.5, ry, 2.2, 0.9, fill=col)
    txt(s, k, 0.55, ry+0.2, 2.1, 0.5, sz=13, bold=True, col=WHITE)
    txt(s, v, 2.8, ry+0.15, 9.8, 0.65, sz=13, col=col if k == "Hach机会" else DARK_GRAY)

# ─── Slide 8: 数据来源 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "数据来源", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

sources = [
    "百度百科（年会基本信息）",
    "中国新闻网（综合大会报道）",
    "深圳政府在线（新闻发布会）",
    "中国日报网（年会预告）",
    "南方新闻网（会议详细报道）",
    "注：以上均为案头研究数据，非一手参会信息",
]
for i, src in enumerate(sources):
    ry = 1.1 + i * 0.85
    bg = RGBColor(250,250,250) if i % 2 == 0 else RGBColor(245,248,250)
    rect(s, 0.5, ry, 12.33, 0.75, fill=bg)
    txt(s, "* " + src, 0.7, ry+0.15, 12, 0.5, sz=13,
        col=LIGHT_GRAY if i == 5 else DARK_GRAY)

rect(s, 0.5, 7.35, 12.33, 0.03, fill=LIGHT_GRAY)
txt(s, "中国城镇供水排水协会2026年会  |  数据来源：案头研究  |  2026-04-22",
    0.5, 7.42, 12.33, 0.3, sz=10, col=LIGHT_GRAY)

out = "/home/agentuser/Work-02_中国水协2026年会.pptx"
prs.save(out)
print("Saved:", out, "(" + str(len(prs.slides)) + " slides)")
