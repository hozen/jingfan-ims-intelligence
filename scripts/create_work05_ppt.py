#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Work-05: 领导力修炼 PPT
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BLUE = RGBColor(0, 126, 181)
DARK_GRAY = RGBColor(77, 77, 77)
LIGHT_GRAY = RGBColor(153, 153, 153)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

def rect(s, x, y, w, h, fill=None):
    shape = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    shape.line.fill.background()
    return shape

def txt(s, text, x, y, w, h, sz=14, bold=False, col=DARK_GRAY,
        align=PP_ALIGN.LEFT):
    txBox = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(sz)
    run.font.bold = bold
    run.font.color.rgb = col
    return txBox

def bullets(s, items, x, y, w, h, sz=13, col=DARK_GRAY):
    txBox = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = item
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(sz)
        run.font.color.rgb = col
    return txBox

# ─── Slide 1: 封面 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 0.12, 7.5, fill=BLUE)
txt(s, "领导力修炼", 0.5, 1.6, 12, 0.9, sz=44, bold=True, col=DARK_GRAY)
txt(s, "情景领导力工作坊", 0.5, 2.55, 12, 0.8, sz=36, bold=True, col=BLUE)
txt(s, "R1 R2 R3 R4  x  S1 S2 S3 S4", 0.5, 3.5, 12, 0.5, sz=20, col=LIGHT_GRAY)
rect(s, 0.5, 4.15, 12.5, 0.03, fill=LIGHT_GRAY)
txt(s, "应用场景：直接下属管理 · 跨部门协作 · 非职权领导力 · 下一代培养",
    0.5, 4.35, 12, 0.4, sz=13, col=DARK_GRAY)
txt(s, "Hach中国IMS产品线  |  2026-04-21", 0.5, 4.85, 12, 0.4, sz=11, col=LIGHT_GRAY)

# ─── Slide 2: 什么是情景领导力 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "什么是情景领导力？", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

rect(s, 0.5, 1.1, 12.33, 1.2, fill=RGBColor(232,244,252))
txt(s, "核心思想：没有最佳的领导风格——只有因人而异的领导风格", 0.7, 1.2, 12, 0.5, sz=18, bold=True, col=BLUE)
txt(s, "同样的行为（如下指令）在不同人身上产生完全不同的效果——取决于对方的准备度", 0.7, 1.7, 12, 0.5, sz=13, col=DARK_GRAY)

principles = [
    "优秀的领导者会根据被领导者的准备度，灵活调整自己的领导风格",
    "准备度 = 能力（知识和技能）x 意愿（动机和自信心）",
    "不存在最好 的领导风格——只有最适合当前这个人当前任务的方式",
    "领导者的职责是帮助被领导者成长，而不是用自己习惯的方式一以贯之",
]
for i, p_text in enumerate(principles):
    ry = 2.5 + i * 1.05
    rect(s, 0.5, ry, 12.33, 0.95, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 0.06, 0.95, fill=BLUE)
    txt(s, str(i+1) + ". " + p_text, 0.7, ry+0.2, 12, 0.6, sz=13, col=DARK_GRAY)

# ─── Slide 3: 四种准备度 R1-R4 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "被领导者的四种准备度", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

readiness = [
    ("R1", "无能力 x 有意愿", "热情的新手", "不知道怎么做，但很想做", RGBColor(180,60,0)),
    ("R2", "无能力 x 无意愿", "消极的学习者", "不知道怎么做，也不想做", RGBColor(150,50,50)),
    ("R3", "有能力 x 无意愿", "能干的犹豫者", "知道怎么做，但不想做", RGBColor(0,100,60)),
    ("R4", "有能力 x 有意愿", "高效的成功者", "知道怎么做，也想做", RGBColor(0,126,181)),
]
for i, (code, formula, name, desc, col) in enumerate(readiness):
    ry = 1.1 + i * 1.55
    rect(s, 0.5, ry, 12.33, 1.4, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 1.5, 1.4, fill=col)
    txt(s, code, 0.55, ry+0.15, 1.4, 0.5, sz=24, bold=True, col=WHITE, align=PP_ALIGN.CENTER)
    txt(s, formula, 0.55, ry+0.65, 1.4, 0.6, sz=11, col=WHITE, align=PP_ALIGN.CENTER)
    txt(s, name, 2.1, ry+0.15, 3.5, 0.5, sz=17, bold=True, col=col)
    txt(s, desc, 2.1, ry+0.7, 10.5, 0.6, sz=13, col=DARK_GRAY)

# ─── Slide 4: 四种领导风格 S1-S4 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "领导者的四种风格", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)
txt(s, "两个维度：任务行为（具体指示） + 关系行为（情感支持）", 0.5, 0.9, 12, 0.4, sz=13, col=LIGHT_GRAY)

styles = [
    ("S1", "指挥式", "高任务 x 低关系",
     "领导者给出明确指令：做什么、怎么做、何时做",
     BLUE, "R1 热情新手"),
    ("S2", "教练式", "高任务 x 高关系",
     "领导者解释决策背后的原因，同时提供情感支持",
     RGBColor(0,100,150), "R2 消极学习者"),
    ("S3", "支持式", "低任务 x 高关系",
     "领导者促进和鼓励，分享想法，邀请提问",
     RGBColor(0,130,100), "R3 能干犹豫者"),
    ("S4", "授权式", "低任务 x 低关系",
     "领导者将责任和决策权交给对方，起到监督作用",
     DARK_GRAY, "R4 高效成功者"),
]
for i, (code, name, dimension, desc, col, target) in enumerate(styles):
    ry = 1.45 + i * 1.45
    rect(s, 0.5, ry, 12.33, 1.3, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 1.5, 1.3, fill=col)
    txt(s, code, 0.55, ry+0.2, 1.4, 0.5, sz=22, bold=True, col=WHITE, align=PP_ALIGN.CENTER)
    txt(s, name, 2.1, ry+0.1, 2.5, 0.5, sz=17, bold=True, col=col)
    txt(s, dimension, 2.1, ry+0.6, 2.5, 0.5, sz=12, col=LIGHT_GRAY)
    txt(s, desc, 4.7, ry+0.2, 6.0, 0.95, sz=13, col=DARK_GRAY)
    rect(s, 10.8, ry+0.25, 2.0, 0.8, fill=RGBColor(240,248,255))
    txt(s, "适用：" + target, 10.85, ry+0.4, 1.9, 0.5, sz=11, col=BLUE)

# ─── Slide 5: R -> S 匹配矩阵 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "R -> S 匹配矩阵", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)
txt(s, "核心原则：领导风格必须与被领导者的准备度匹配", 0.5, 0.9, 12, 0.4, sz=13, col=LIGHT_GRAY)

col_ws = [1.6, 3.2, 2.8, 1.2, 4.5]
col_xs = [0.5]
for w in col_ws[:-1]:
    col_xs.append(col_xs[-1] + w)

headers = ["准备度", "描述", "领导者风格", "S代码", "关键行为"]
for ci, (h, cx, cw) in enumerate(zip(headers, col_xs, col_ws)):
    rect(s, cx, 1.4, cw, 0.45, fill=BLUE)
    txt(s, h, cx+0.08, 1.42, cw-0.1, 0.4, sz=12, bold=True, col=WHITE)

rows = [
    ("R1", "热情新手", "S1 指挥式", "S1", "给指令、告诉做什么和怎么做"),
    ("R2", "消极学习者", "S2 教练式", "S2", "解释原因、提供方向、支持鼓励"),
    ("R3", "能干犹豫者", "S3 支持式", "S3", "分享想法、邀请参与、共同决策"),
    ("R4", "高效成功者", "S4 授权式", "S4", "把责任和决策权交给对方"),
]
row_cols = [BLUE, RGBColor(0,100,150), RGBColor(0,130,100), DARK_GRAY]
for ri, (r, desc, style, scode, behavior) in enumerate(rows):
    ry = 1.88 + ri * 1.2
    bg = RGBColor(250,250,250) if ri % 2 == 0 else RGBColor(245,248,250)
    rect(s, 0.5, ry, 12.33, 1.1, fill=bg)
    vals = [r, desc, style, scode, behavior]
    for ci, (val, cx, cw) in enumerate(zip(vals, col_xs, col_ws)):
        col = row_cols[ri] if ci in (0,2,3) else DARK_GRAY
        is_code = (ci == 3)
        txt(s, val, cx+0.08, ry+0.25, cw-0.1, 0.6,
            sz=14 if is_code else 13,
            bold=is_code, col=col)

rect(s, 0.5, 6.75, 12.33, 0.65, fill=RGBColor(255,245,240))
rect(s, 0.5, 6.75, 0.06, 0.65, fill=RGBColor(180,60,0))
txt(s, "WARNING: 错配代价 - S1用于R3会被感到微观管理；S4用于R1会导致缺乏指导焦虑增加",
    0.7, 6.82, 12, 0.5, sz=12, col=RGBColor(180,60,0))

# ─── Slide 6: 应用场景 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "应用场景", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

scenarios = [
    ("直接下属管理", "有职权",
     "根据下属的R层级，灵活切换S1-S4风格",
     "绩效谈话、任务分配、新人带教"),
    ("跨部门协作", "无职权",
     "识别对方的准备度，用对方能接受的方式影响",
     "项目推动、会议协调、资源争取"),
    ("销售团队管理", "有职权",
     "S1带新人拿单 -> S3带资深客户攻关 -> S4授权区域负责人",
     "客户拜访管理、区域授权"),
    ("下一代培养", "无职权（家长角色）",
     "孩子在不同成长阶段需要不同领导风格",
     "作业/兴趣引导、独立性培养"),
]
for i, (scene, role, principle, examples) in enumerate(scenarios):
    ry = 1.05 + i * 1.55
    rect(s, 0.5, ry, 12.33, 1.4, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 2.5, 1.4, fill=BLUE)
    txt(s, scene, 0.6, ry+0.2, 2.3, 0.5, sz=14, bold=True, col=WHITE)
    txt(s, role, 0.6, ry+0.7, 2.3, 0.5, sz=12, col=RGBColor(200,230,250))
    rect(s, 3.1, ry+0.1, 4.3, 1.2, fill=RGBColor(245,248,250))
    txt(s, principle, 3.2, ry+0.35, 4.1, 0.8, sz=12, col=DARK_GRAY)
    txt(s, "例如：" + examples, 7.5, ry+0.35, 5.5, 0.8, sz=12, col=LIGHT_GRAY)

# ─── Slide 7: 成长型 vs 固定型思维 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "思维模式：成长型 vs 固定型", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)
txt(s, "领导风格之外：思维模式决定一个人如何接收反馈和面对挑战", 0.5, 0.9, 12, 0.4, sz=13, col=LIGHT_GRAY)

mindset = [
    ("固定型思维", "能力是固定的 —— 要证明自己",
     "被批评 -> 感到被攻击；遇到困难 -> 放弃",
     RGBColor(180,60,0)),
    ("成长型思维", "能力是可以发展的 —— 要学习",
     "被批评 -> 获得反馈；遇到困难 -> 找方法",
     RGBColor(0,126,100)),
]
for i, (name, belief, behavior, col) in enumerate(mindset):
    ry = 1.5 + i * 2.0
    rect(s, 0.5, ry, 12.33, 1.8, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 0.06, 1.8, fill=col)
    rect(s, 0.65, ry+0.1, 3.2, 1.6, fill=col)
    txt(s, name, 0.7, ry+0.55, 3.0, 0.5, sz=16, bold=True, col=WHITE)
    txt(s, "信念：" + belief, 4.0, ry+0.2, 8.5, 0.6, sz=13, col=DARK_GRAY)
    txt(s, "面对挑战：" + behavior, 4.0, ry+0.9, 8.5, 0.6, sz=13, col=col)

# ─── Slide 8: 核心要点 ───
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 13.33, 0.08, fill=BLUE)
txt(s, "核心要点", 0.5, 0.35, 12, 0.6, sz=28, bold=True, col=DARK_GRAY)

takeaways = [
    ("灵活切换风格", "没有最好的领导风格——只有适合当前这个人、当前任务的风格"),
    ("识别准备度", "先判断对方是R1-R4中的哪一个，再选择对应S1-S4风格"),
    ("错配要纠正", "长期用错误风格会导致对方准备度退化（S1用于R3=信心受损）"),
    ("思维模式影响沟通", "固定型思维者需要先建立安全感；成长型思维者可以直接给反馈"),
    ("实践才能内化", "案例练习是掌握情景领导力的唯一路径——直到形成本能反应"),
]
for i, (title, body) in enumerate(takeaways):
    ry = 1.1 + i * 1.2
    rect(s, 0.5, ry, 12.33, 1.1, fill=RGBColor(250,250,250))
    rect(s, 0.5, ry, 0.06, 1.1, fill=BLUE)
    rect(s, 0.65, ry+0.15, 0.55, 0.55, fill=BLUE)
    txt(s, str(i+1), 0.65, ry+0.15, 0.55, 0.55, sz=18, bold=True, col=WHITE, align=PP_ALIGN.CENTER)
    txt(s, title, 1.35, ry+0.12, 4.0, 0.5, sz=15, bold=True, col=DARK_GRAY)
    txt(s, body, 1.35, ry+0.6, 11.3, 0.45, sz=12, col=DARK_GRAY)

rect(s, 0.5, 7.35, 12.33, 0.03, fill=LIGHT_GRAY)
txt(s, "情景领导力  |  2026-04-21  |  Hach中国IMS产品线", 0.5, 7.42, 12.33, 0.3, sz=10, col=LIGHT_GRAY)

out = "/home/agentuser/Work-05_领导力修炼.pptx"
prs.save(out)
print("Saved:", out, "(" + str(len(prs.slides)) + " slides)")
